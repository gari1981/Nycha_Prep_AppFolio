from __future__ import annotations
import argparse
import logging
import os
import re
import time
from dataclasses import dataclass
from datetime import date, datetime, time as datetime_time
from pathlib import Path
from typing import Any
APPFOLIO_URL = os.getenv('APPFOLIO_URL', '')
GOOGLE_SCOPES = ('https://www.googleapis.com/auth/spreadsheets',)
SEARCH_RESULT_TIMEOUT_SECONDS = 30
APPFOLIO_SEARCH_RETRY_STEPS = ((100, 12, 2), (250, 30, 5), (400, 60, 0))
GOOGLE_RETRY_DELAYS_SECONDS = (1, 2, 4, 8, 15, 30, 45)

@dataclass(frozen=True)
class ReportJob:
    report_name: str
    destination_sheet: str

def require_dependencies() -> None:
    missing: list[str] = []
    for import_name, package_name in (('googleapiclient.discovery', 'google-api-python-client'), ('google.oauth2.credentials', 'google-auth'), ('google_auth_oauthlib.flow', 'google-auth-oauthlib'), ('openpyxl', 'openpyxl'), ('playwright.sync_api', 'playwright')):
        try:
            __import__(import_name)
        except ModuleNotFoundError:
            missing.append(package_name)
    if missing:
        raise RuntimeError('Missing Python packages: ' + ', '.join(sorted(set(missing))))

def retryable_connectivity_error(exc: Exception) -> bool:
    """Return whether a Google request failed for a transient network reason."""
    if isinstance(exc, (TimeoutError, ConnectionError, OSError)):
        return True
    status = getattr(getattr(exc, 'resp', None), 'status', None)
    if status in {408, 425, 429} or (isinstance(status, int) and status >= 500):
        return True
    message = str(exc).casefold()
    return any((marker in message for marker in ('timed out', 'timeout', 'connection reset', 'connection aborted', 'connection refused', 'remote end closed', 'unexpected_eof', 'unexpected eof', 'temporary failure', 'name resolution', 'servernotfound', 'service unavailable', 'bad gateway', 'rate limit', 'ssl')))

def browser_target_closed_error(exc: Exception) -> bool:
    """Return whether Playwright lost the browser process or its active target."""
    message = f'{type(exc).__name__}: {exc}'.casefold()
    return any((marker in message for marker in ('targetclosederror', 'target page, context or browser has been closed', 'browser has been closed', 'context has been closed')))

def execute_google(request, attempts: int | None=None):
    max_attempts = attempts or len(GOOGLE_RETRY_DELAYS_SECONDS) + 1
    for attempt in range(1, max_attempts + 1):
        try:
            return request.execute(num_retries=3)
        except Exception as exc:
            if attempt == max_attempts or not retryable_connectivity_error(exc):
                raise
            delay = GOOGLE_RETRY_DELAYS_SECONDS[min(attempt - 1, len(GOOGLE_RETRY_DELAYS_SECONDS) - 1)]
            logging.warning('RETRY Google Sheets connectivity failure on attempt %s/%s; waiting %s second(s): %s', attempt, max_attempts, delay, exc)
            time.sleep(delay)

def build_sheets_service(args: argparse.Namespace):
    import google.auth
    import google_auth_httplib2
    import httplib2
    from google.auth.transport.requests import Request
    from google.oauth2.credentials import Credentials
    from google_auth_oauthlib.flow import InstalledAppFlow
    from googleapiclient.discovery import build
    credentials = None
    credentials_path = Path(args.google_credentials_file).expanduser() if args.google_credentials_file else None
    oauth_client_path = Path(args.google_oauth_client_file).expanduser() if args.google_oauth_client_file else None
    token_path = Path(args.google_token_file).expanduser()
    if credentials_path:
        credentials, _ = google.auth.load_credentials_from_file(str(credentials_path), scopes=list(GOOGLE_SCOPES))
    elif token_path.is_file():
        credentials = Credentials.from_authorized_user_file(str(token_path), scopes=list(GOOGLE_SCOPES))
    if credentials and credentials.expired and credentials.refresh_token:
        for attempt in range(1, len(GOOGLE_RETRY_DELAYS_SECONDS) + 2):
            try:
                credentials.refresh(Request())
                break
            except Exception as exc:
                if attempt == len(GOOGLE_RETRY_DELAYS_SECONDS) + 1 or not retryable_connectivity_error(exc):
                    raise
                delay = GOOGLE_RETRY_DELAYS_SECONDS[attempt - 1]
                logging.warning('RETRY Google authorization refresh attempt %s failed; waiting %s second(s): %s', attempt, delay, exc)
                time.sleep(delay)
    if not credentials or not credentials.valid:
        if not oauth_client_path or not oauth_client_path.is_file():
            raise RuntimeError('Google authorization is unavailable. Pass an OAuth client file once to create the token.')
        if args.headless:
            raise RuntimeError('First Google OAuth approval must be completed without --headless.')
        flow = InstalledAppFlow.from_client_secrets_file(str(oauth_client_path), scopes=list(GOOGLE_SCOPES))
        credentials = flow.run_local_server(port=0, open_browser=True)
        token_path.parent.mkdir(parents=True, exist_ok=True)
        token_path.write_text(credentials.to_json(), encoding='utf-8')
        token_path.chmod(384)
    http = google_auth_httplib2.AuthorizedHttp(credentials, http=httplib2.Http(timeout=max(30, args.google_timeout_seconds)))
    return build('sheets', 'v4', http=http, cache_discovery=False)

def quote_sheet_name(name: str) -> str:
    return "'" + name.replace("'", "''") + "'"

def wait_for_appfolio_login(page, args: argparse.Namespace) -> None:
    logging.info("STARTED Opening AppFolio.")

    def first_visible(*locators):
        for locator in locators:
            if locator is None:
                continue
            try:
                candidate = locator.first
                if candidate.count() > 0 and candidate.is_visible():
                    return candidate
            except Exception:
                continue
        return None

    def login_controls():
        """Resolve both the former and current AppFolio login markup."""
        try:
            email_by_label = page.get_by_label("Email", exact=True)
        except Exception:
            email_by_label = None
        try:
            email_by_role = page.get_by_role("textbox", name="Email", exact=True)
        except Exception:
            email_by_role = None
        try:
            email_by_css = page.locator(
                "input[type='email'], input[name='username'], input[name='email']"
            )
        except Exception:
            email_by_css = None

        try:
            password_by_label = page.get_by_label("Password", exact=True)
        except Exception:
            password_by_label = None
        try:
            password_by_css = page.locator("input[type='password']")
        except Exception:
            password_by_css = None

        try:
            login_by_role = page.get_by_role(
                "button", name=re.compile(r"^(?:log|sign)\s*in$", re.I)
            )
        except Exception:
            login_by_role = None
        try:
            login_by_css = page.locator("button[type='submit'], input[type='submit']")
        except Exception:
            login_by_css = None

        try:
            email_by_test_id = page.get_by_test_id("email-input")
        except Exception:
            email_by_test_id = None
        try:
            password_by_test_id = page.get_by_test_id("password-input")
        except Exception:
            password_by_test_id = None
        try:
            login_by_test_id = page.get_by_test_id("sign-in-button")
        except Exception:
            login_by_test_id = None

        email = first_visible(email_by_test_id, email_by_label, email_by_role, email_by_css)
        password = first_visible(password_by_test_id, password_by_label, password_by_css)
        submit = first_visible(login_by_test_id, login_by_role, login_by_css)
        return (email, password, submit) if email and password and submit else None

    def logged_in() -> bool:
        try:
            search = page.get_by_role("searchbox", name="Search")
            return search.count() > 0 and search.first.is_visible()
        except Exception:
            return False

    from appfolio_sms import SmsVerification
    service = getattr(args, 'sms_sheets_service', None)
    sms = SmsVerification(service) if service is not None else None
    manual = not args.headless and not args.non_interactive_login
    page.goto(args.appfolio_url, wait_until="domcontentloaded", timeout=120000)
    deadline = time.monotonic() + 240
    submitted = False
    logging.info("WAITING for AppFolio login; SMS verification automation is %s.",
                 "enabled" if sms else "unavailable")
    while not logged_in():
        if page.is_closed():
            raise RuntimeError("AppFolio login window was closed before login completed.")
        if not manual and time.monotonic() >= deadline:
            raise RuntimeError("AppFolio login did not complete within 240 seconds. Check credentials or verification requirements.")
        if sms is not None:
            try:
                sms.poll(page)
            except Exception as exc:
                if browser_target_closed_error(exc):
                    raise
                if not manual:
                    raise RuntimeError("Automatic SMS verification failed (" + type(exc).__name__ +
                                       "). Check SMS sheet access, message arrival and verification form.") from None
                logging.warning("SMS automation unavailable (%s); browser remains open for manual verification.", type(exc).__name__)
                sms = None
        if not submitted and args.appfolio_email and args.appfolio_password:
            controls = login_controls()
            if controls:
                email_input, password_input, login_button = controls
                if (email_input.is_enabled() and password_input.is_enabled()
                        and login_button.is_enabled()):
                    email_input.fill(args.appfolio_email)
                    password_input.fill(args.appfolio_password)
                    login_button.click(no_wait_after=True)
                    submitted = True
                    logging.info("SUBMITTED AppFolio username/password; waiting for login or verification.")
        page.wait_for_timeout(1000)
    logging.info("SUCCESS AppFolio login completed.")



def safe_filename(value: str) -> str:
    return re.sub('[^A-Za-z0-9._-]+', '_', value).strip('_')

def appfolio_search_input_value(search) -> str | None:
    """Best-effort readback used to detect React dropping a typed query."""
    try:
        return str(search.input_value(timeout=5000))
    except AttributeError:
        value = getattr(search, 'value', None)
        return None if value is None else str(value)
    except Exception:
        try:
            value = search.get_attribute('value')
            return None if value is None else str(value)
        except Exception:
            return None

def progressive_appfolio_search(page, name: str, result_waiter, result_kind: str):
    """Try a search quickly, then retype it progressively slower when needed."""
    total = len(APPFOLIO_SEARCH_RETRY_STEPS)
    for attempt, (delay_ms, result_wait_seconds, retry_pause_seconds) in enumerate(APPFOLIO_SEARCH_RETRY_STEPS, start=1):
        try:
            page.bring_to_front()
            search = page.get_by_role('searchbox', name='Search')
            search.click()
            search.fill('')
            page.wait_for_timeout(250 * attempt)
            search.press_sequentially(name, delay=delay_ms)
            page.wait_for_timeout(750 * attempt)
            actual = appfolio_search_input_value(search)
            if actual is not None and actual.strip() != name:
                logging.warning('RETRY AppFolio dropped the %s search text on attempt %s/%s (expected %r, found %r).', result_kind, attempt, total, name, actual)
            else:
                logging.info('Typed AppFolio %s search on attempt %s/%s at %sms per character: %s', result_kind, attempt, total, delay_ms, name)
            result = result_waiter(page, name, timeout_seconds=result_wait_seconds)
            if result is not None:
                if attempt > 1:
                    logging.info('SUCCESS AppFolio %s search recovered on attempt %s/%s.', result_kind, attempt, total)
                return result
        except Exception as exc:
            logging.warning('RETRY AppFolio %s search attempt %s/%s failed: %s', result_kind, attempt, total, exc)
        if attempt < total:
            logging.warning('RETRY Repeating AppFolio %s search more slowly after %s second(s): %s', result_kind, retry_pause_seconds, name)
            page.wait_for_timeout(retry_pause_seconds * 1000)
    return None

def capture_search_diagnostics(page, name: str, diagnostics_dir: Path) -> list[Path]:
    """Save enough local evidence to diagnose a changed or slow AppFolio result UI."""
    diagnostics_dir.mkdir(parents=True, exist_ok=True)
    stem = f"{datetime.now():%Y%m%d_%H%M%S}_{safe_filename(name) or 'search'}"
    saved: list[Path] = []
    screenshot_path = diagnostics_dir / f'{stem}.png'
    try:
        page.screenshot(path=str(screenshot_path), full_page=True)
        saved.append(screenshot_path)
    except Exception as exc:
        logging.warning('Could not capture AppFolio search screenshot: %s', exc)
    html_path = diagnostics_dir / f'{stem}.html'
    try:
        html_path.write_text(page.content(), encoding='utf-8')
        saved.append(html_path)
    except Exception as exc:
        logging.warning('Could not capture AppFolio search HTML: %s', exc)
    details_path = diagnostics_dir / f'{stem}.txt'
    try:
        details_path.write_text('\n'.join([f"captured_at={datetime.now().astimezone().isoformat(timespec='seconds')}", f'search_value={name}', f'url={page.url}', f'title={page.title()}']) + '\n', encoding='utf-8')
        saved.append(details_path)
    except Exception as exc:
        logging.warning('Could not capture AppFolio search details: %s', exc)
    if saved:
        logging.error('Saved AppFolio search diagnostics: %s', ', '.join((str(path) for path in saved)))
    return saved

def first_visible(locator):
    """Return the first currently visible match without trusting DOM order."""
    for index in range(min(locator.count(), 50)):
        candidate = locator.nth(index)
        if candidate.is_visible():
            return candidate
    return None

def report_page_scopes(page) -> list[Any]:
    scopes = [page]
    main_frame = getattr(page, 'main_frame', None)
    for frame in getattr(page, 'frames', []):
        if frame is not main_frame:
            scopes.append(frame)
    return scopes

def capture_report_diagnostics(page, report_name: str, diagnostics_dir: Path, stage: str) -> list[Path]:
    saved = capture_search_diagnostics(page, f'report_{stage}_{report_name}', diagnostics_dir)
    controls_path = diagnostics_dir / f'{datetime.now():%Y%m%d_%H%M%S}_report_{safe_filename(stage)}_{safe_filename(report_name)}_controls.txt'
    lines = []
    for scope_index, scope in enumerate(report_page_scopes(page)):
        lines.append(f"scope={scope_index} url={getattr(scope, 'url', '')}")
        try:
            controls = scope.locator("button, a, [role='button'], [role='link'], [role='menuitem']")
            for index in range(min(controls.count(), 500)):
                control = controls.nth(index)
                if not control.is_visible():
                    continue
                text = ' '.join(control.inner_text().split())
                aria = control.get_attribute('aria-label') or ''
                testid = control.get_attribute('data-testid') or ''
                lines.append(f'  text={text!r} aria={aria!r} testid={testid!r}')
        except Exception as exc:
            lines.append(f'  inspection_error={exc}')
    try:
        controls_path.write_text('\n'.join(lines) + '\n', encoding='utf-8')
        saved.append(controls_path)
        logging.error('Saved report-control diagnostics: %s', controls_path)
    except Exception as exc:
        logging.warning('Could not save report-control diagnostics: %s', exc)
    return saved

def wait_for_saved_report_link(page, report_name: str, timeout_seconds: float=SEARCH_RESULT_TIMEOUT_SECONDS):
    deadline = time.monotonic() + timeout_seconds
    while time.monotonic() < deadline:
        for link in (page.get_by_role('link', name=report_name, exact=True), page.get_by_role('link', name=re.compile(re.escape(report_name), re.I))):
            try:
                candidate = first_visible(link)
                if candidate is not None:
                    return candidate
            except Exception:
                pass
        page.wait_for_timeout(250)
    return None

def open_saved_report(page, report_name: str, diagnostics_dir: Path):
    link = progressive_appfolio_search(page, report_name, wait_for_saved_report_link, 'saved-report')
    if link is None:
        saved = capture_report_diagnostics(page, report_name, diagnostics_dir, 'search_result_missing')
        suffix = f" Diagnostics: {', '.join((str(path) for path in saved))}" if saved else ''
        raise RuntimeError(f'Saved AppFolio report link was not found: {report_name}.{suffix}')
    try:
        with page.expect_popup(timeout=90000) as popup_info:
            link.click(no_wait_after=True)
        report_page = popup_info.value
    except Exception as exc:
        saved = capture_report_diagnostics(page, report_name, diagnostics_dir, 'popup_missing')
        suffix = f" Diagnostics: {', '.join((str(path) for path in saved))}" if saved else ''
        raise RuntimeError(f'Saved AppFolio report did not open in its popup: {report_name}.{suffix}') from exc
    try:
        report_page.wait_for_load_state('domcontentloaded', timeout=120000)
    except Exception as exc:
        logging.warning('RETRY Report popup DOMContentLoaded wait expired; continuing with the Actions-control readiness check: %s', exc)
    report_page.wait_for_timeout(3000)
    report_page.bring_to_front()
    last_error = None
    for ready_attempt, ready_timeout in enumerate((60000, 180000), start=1):
        actions = report_page.get_by_role('button', name='Actions')
        try:
            actions.wait_for(state='visible', timeout=ready_timeout)
            break
        except Exception as exc:
            last_error = exc
            if ready_attempt == 1:
                logging.warning('RETRY Report %s did not become ready quickly; reloading once and allowing three minutes.', report_name)
                try:
                    report_page.reload(wait_until='domcontentloaded', timeout=120000)
                except Exception as reload_exc:
                    logging.warning('Report reload did not finish cleanly; readiness polling will continue: %s', reload_exc)
    else:
        saved = capture_report_diagnostics(report_page, report_name, diagnostics_dir, 'ready')
        suffix = f" Diagnostics: {', '.join((str(path) for path in saved))}" if saved else ''
        raise RuntimeError(f'AppFolio report did not expose an Actions control after a reload: {report_name}.{suffix}') from last_error
    logging.info('SUCCESS Report is ready: %s', report_name)
    return report_page

def download_report(
    report_page, report_name: str, download_dir: Path, diagnostics_dir: Path
) -> Path:
    # Log lifecycle events without logging URLs, credentials or report content.
    context = report_page.context
    browser = context.browser
    report_page.on("close", lambda *_: logging.warning("EXPORT EVENT: report tab closed"))
    report_page.on("crash", lambda *_: logging.error("EXPORT EVENT: report tab crashed"))
    context.on("close", lambda *_: logging.warning("EXPORT EVENT: browser context closed"))
    if browser is not None:
        browser.on("disconnected", lambda *_: logging.warning("EXPORT EVENT: browser disconnected"))
    last_error = None
    for attempt, control_timeout in enumerate((30_000, 60_000, 120_000), start=1):
        try:
            export_excel = report_page.get_by_text("Export as Excel", exact=True)
            if first_visible(export_excel) is None:
                report_page.get_by_role("button", name="Actions").click(
                    no_wait_after=True
                )
                report_page.wait_for_timeout(750 * attempt)
            export_excel.wait_for(state="visible", timeout=control_timeout)
            with report_page.expect_download(
                timeout=max(120_000, control_timeout)
            ) as download_info:
                export_excel.click(no_wait_after=True)
            target = download_dir / (
                f"{safe_filename(report_name)}_{datetime.now():%Y%m%d_%H%M%S}.xlsx"
            )
            download = download_info.value
            logging.info("STARTED Waiting for Excel download to complete: %s", report_name)
            temporary_path = download.path()
            if temporary_path is None:
                raise RuntimeError("Excel download returned no local file.")
            download.save_as(target)
            if not target.is_file() or target.stat().st_size == 0:
                raise RuntimeError(f"Downloaded report is missing or empty: {target}")
            if attempt > 1:
                logging.info(
                    "SUCCESS Report download recovered on attempt %s/3: %s",
                    attempt,
                    report_name,
                )
            logging.info("SUCCESS Downloaded %s", target.name)
            return target
        except Exception as exc:
            last_error = exc
            page_closed = report_page.is_closed()
            browser_connected = browser.is_connected() if browser is not None else None
            logging.error(
                "EXPORT FAILURE: stage=download report_tab_closed=%s browser_connected=%s error_type=%s",
                page_closed, browser_connected, type(exc).__name__,
            )
            if browser_target_closed_error(exc) or page_closed or browser_connected is False:
                # Preserve Download.save_as failure instead of masking it by
                # calling wait_for_timeout on a closed page.
                raise
            if attempt < 3:
                pause_seconds = 3 * attempt
                logging.warning(
                    "RETRY Report download attempt %s/3 failed; waiting %s "
                    "second(s) and trying more slowly: %s",
                    attempt,
                    pause_seconds,
                    exc,
                )
                report_page.wait_for_timeout(pause_seconds * 1000)
    capture_report_diagnostics(
        report_page, report_name, diagnostics_dir, "download_failed"
    )
    raise RuntimeError(
        f"Could not download {report_name} after three progressively slower attempts."
    ) from last_error

def google_value(value: Any) -> Any:
    if value is None:
        return ''
    if isinstance(value, datetime):
        return value.isoformat(sep=' ', timespec='seconds')
    if isinstance(value, (date, datetime_time)):
        return value.isoformat()
    if isinstance(value, (str, int, float, bool)):
        return value
    return str(value)

def read_excel_values(path: Path) -> list[list[Any]]:
    from openpyxl import load_workbook
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        visible = [sheet for sheet in workbook.worksheets if sheet.sheet_state == 'visible']
        if not visible:
            raise RuntimeError(f'No visible worksheet found in {path.name}')
        rows = [[google_value(value) for value in row] for row in visible[0].iter_rows(values_only=True)]
    finally:
        workbook.close()
    while rows and (not any((value != '' for value in rows[-1]))):
        rows.pop()
    if not rows:
        raise RuntimeError(f'The first visible worksheet in {path.name} contains no values.')
    logging.info('SUCCESS Prepared %s rows from %s', len(rows), path.name)
    return rows

def ensure_sheet_size(service, spreadsheet_id: str, sheet_name: str, rows: int, columns: int) -> None:
    response = execute_google(service.spreadsheets().get(spreadsheetId=spreadsheet_id, fields='sheets(properties(sheetId,title,gridProperties(rowCount,columnCount)))'))
    props = next((sheet['properties'] for sheet in response.get('sheets', []) if sheet['properties']['title'] == sheet_name), None)
    if not props:
        raise RuntimeError(f'Destination tab not found: {sheet_name}')
    grid = props.get('gridProperties', {})
    target_rows = max(int(grid.get('rowCount', 1)), rows, 1)
    target_columns = max(int(grid.get('columnCount', 1)), columns, 1)
    if target_rows == grid.get('rowCount') and target_columns == grid.get('columnCount'):
        return
    execute_google(service.spreadsheets().batchUpdate(spreadsheetId=spreadsheet_id, body={'requests': [{'updateSheetProperties': {'properties': {'sheetId': props['sheetId'], 'gridProperties': {'rowCount': target_rows, 'columnCount': target_columns}}, 'fields': 'gridProperties(rowCount,columnCount)'}}]}))

def replace_sheet_values(service, spreadsheet_id: str, sheet_name: str, rows: list[list[Any]]) -> None:
    width = max((len(row) for row in rows), default=1)
    ensure_sheet_size(service, spreadsheet_id, sheet_name, len(rows), width)
    quoted = quote_sheet_name(sheet_name)
    values_api = service.spreadsheets().values()
    execute_google(values_api.clear(spreadsheetId=spreadsheet_id, range=quoted, body={}))
    for start in range(0, len(rows), 5000):
        execute_google(values_api.update(spreadsheetId=spreadsheet_id, range=f'{quoted}!A{start + 1}', valueInputOption='RAW', body={'majorDimension': 'ROWS', 'values': rows[start:start + 5000]}))
    logging.info('SUCCESS Replaced values in %s', sheet_name)
